from __future__ import annotations
import uuid

from fastapi import APIRouter, Depends, HTTPException, status
from sqlalchemy import select
from sqlalchemy.dialects.postgresql import insert
from sqlalchemy.orm import Session

from app.database import get_db
from app.models.ai_mention import AiMention
from app.models.client import Client
from app.models.enums import AiMentionSource
from app.schemas.ai_mention import AIMentionManualCreate, AIMentionResponse

from app.dependencies import RequireRole
from app.models.enums import UserRole

router = APIRouter(
    prefix="/clients/{client_id}/ai_mentions",
    tags=["ai_mentions"],
    dependencies=[Depends(RequireRole([UserRole.super_admin, UserRole.manager, UserRole.user]))]
)

from datetime import date

@router.get("")
def get_ai_mentions(
    client_id: uuid.UUID,
    page: int = 1,
    page_size: int = 25,
    start_date: date | None = None,
    end_date: date | None = None,
    db: Session = Depends(get_db),
) -> dict:
    from sqlalchemy import func
    query = select(AiMention).where(
        AiMention.client_id == client_id,
        AiMention.source == AiMentionSource.manual
    )
    if start_date:
        query = query.where(AiMention.month >= start_date)
    if end_date:
        query = query.where(AiMention.month <= end_date)
        
    total = db.execute(select(func.count()).select_from(query.subquery())).scalar() or 0
    mentions = db.execute(query.order_by(AiMention.captured_on.desc()).offset((page - 1) * page_size).limit(page_size)).scalars().all()
    
    records = []
    for m in mentions:
        records.append({
            "id": str(m.id),
            "captured_on": m.captured_on.isoformat(),
            "platform": m.platform.value if m.platform else None,
            "prompt": m.prompt.prompt_text if m.prompt else None,
            "mentioned": m.mentioned,
            "cited_pages": m.cited_pages
        })
        
    return {"items": records, "total": total, "page": page, "page_size": page_size}

@router.post("/manual", response_model=AIMentionResponse, status_code=status.HTTP_201_CREATED)
def create_manual_ai_mention(
    client_id: uuid.UUID, mention_in: AIMentionManualCreate, db: Session = Depends(get_db)
):
    """Manually enter an AI mention record for a client."""
    client = db.execute(select(Client).where(Client.id == client_id)).scalar_one_or_none()
    if not client:
        raise HTTPException(status_code=404, detail="Client not found")

    stmt = insert(AiMention).values(
        client_id=client_id,
        month=mention_in.captured_on.replace(day=1),
        prompt_id=mention_in.prompt_id,
        platform=mention_in.platform,
        captured_on=mention_in.captured_on,
        mentioned=mention_in.mentioned,
        cited_pages=[page.model_dump() for page in mention_in.cited_pages] if mention_in.cited_pages else None,
        source=AiMentionSource.manual,
        raw_response=None
    )
    stmt = stmt.on_conflict_do_update(
        constraint="uq_aimention_client_month_prompt_plat",
        set_={
            "mentioned": stmt.excluded.mentioned,
            "cited_pages": stmt.excluded.cited_pages,
            "captured_on": stmt.excluded.captured_on
        }
    )
    result = db.execute(stmt.returning(AiMention))
    mention = result.scalar_one()
    db.commit()
    return mention

@router.post("/bulk", status_code=status.HTTP_201_CREATED)
def create_bulk_ai_mentions(
    client_id: uuid.UUID, mentions_in: list[AIMentionManualCreate], db: Session = Depends(get_db)
):
    """Manually enter multiple AI mention records via bulk upload."""
    client = db.execute(select(Client).where(Client.id == client_id)).scalar_one_or_none()
    if not client:
        raise HTTPException(status_code=404, detail="Client not found")

    mentions = []
    for mention_in in mentions_in:
        mentions.append({
            "client_id": client_id,
            "month": mention_in.captured_on.replace(day=1),
            "prompt_id": mention_in.prompt_id,
            "platform": mention_in.platform,
            "captured_on": mention_in.captured_on,
            "mentioned": mention_in.mentioned,
            "cited_pages": [page.model_dump() for page in mention_in.cited_pages] if mention_in.cited_pages else None,
            "source": AiMentionSource.manual,
            "raw_response": None
        })

    if mentions:
        stmt = insert(AiMention).values(mentions)
        stmt = stmt.on_conflict_do_update(
            constraint="uq_aimention_client_month_prompt_plat",
            set_={
                "mentioned": stmt.excluded.mentioned,
                "cited_pages": stmt.excluded.cited_pages,
                "captured_on": stmt.excluded.captured_on
            }
        )
        db.execute(stmt)
        db.commit()

    return {"status": "success", "rows_inserted": len(mentions)}

import csv
import io
from fastapi import File, UploadFile
from datetime import datetime

@router.post("/upload_csv", status_code=201)
def upload_ai_mentions_csv(
    client_id: uuid.UUID,
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
) -> dict:
    if not (file.filename.endswith(".csv") or file.filename.endswith(".xlsx")):
        raise HTTPException(status_code=400, detail="Must be a CSV or Excel file")
    
    rows = []
    raw_headers = []
    
    if file.filename.endswith(".xlsx"):
        import openpyxl
        content = file.file.read()
        wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
        ws = wb.active
        if not ws or ws.max_row < 1:
            raise HTTPException(status_code=400, detail="Empty Excel file")
        raw_headers = [str(c.value).strip() if c.value else "" for c in ws[1]]
        for row in ws.iter_rows(min_row=2, values_only=True):
            row_dict = {}
            for k, v in zip(raw_headers, row):
                row_dict[k] = v
            rows.append(row_dict)
    else:
        content = file.file.read().decode("utf-8")
        reader = csv.reader(io.StringIO(content))
        header_row = next(reader, None)
        if not header_row:
            raise HTTPException(status_code=400, detail="Empty CSV")
        raw_headers = [c.strip() for c in header_row]
        for row in reader:
            row_dict = {}
            for k, v in zip(raw_headers, row):
                row_dict[k] = v
            rows.append(row_dict)
        
    normalized_headers = [h.lower().replace(" ", "_") for h in raw_headers]
    
    if "prompt" not in normalized_headers and "prompts" not in normalized_headers:
        raise HTTPException(status_code=400, detail="CSV/Excel must contain 'Prompts' column")
        
    from app.models.enums import AiPlatform
    platform_map = {}
    for h in raw_headers:
        norm_h = h.lower().replace(" ", "")
        if norm_h in ["chatgpt", "gpt"]:
            platform_map[h] = AiPlatform.chatgpt
        elif norm_h in ["claude"]:
            platform_map[h] = AiPlatform.claude
        elif norm_h in ["gemini", "googlegemini"]:
            platform_map[h] = AiPlatform.gemini
        elif norm_h in ["perplexity"]:
            platform_map[h] = AiPlatform.perplexity
        elif norm_h in ["grok"]:
            platform_map[h] = AiPlatform.grok
        elif norm_h in ["aioverview", "googleaioverview"]:
            platform_map[h] = AiPlatform.google_ai_overview

    mentions_to_insert = []
    success_count = 0
    errors = []
    
    target_month = datetime.today().date().replace(day=1)
    
    for row_num, row in enumerate(rows, start=2):
        row_norm = {k.lower().replace(" ", "_"): v for k, v in row.items() if k}
        try:
            prompt = str(row_norm.get("prompt") or row_norm.get("prompts") or "").strip()
            if not prompt or prompt == "None":
                continue
                
            month = target_month
            month_str = str(row_norm.get("month") or row_norm.get("date") or "").strip()
            if month_str:
                for fmt in ("%b'%y", "%B %Y", "%Y-%m-%d", "%m/%d/%Y", "%m/%d/%y", "%Y/%m/%d"):
                    try:
                        clean_str = month_str.split(" ")[0].split("T")[0] if 'T' in month_str else month_str
                        month = datetime.strptime(clean_str, fmt).date().replace(day=1)
                        break
                    except ValueError:
                        pass
                
            from app.models.ai_prompt import AiPrompt
            prompt_obj = db.execute(
                select(AiPrompt).where(
                    AiPrompt.client_id == client_id,
                    AiPrompt.prompt_text == prompt
                )
            ).scalar_one_or_none()
            
            if not prompt_obj:
                prompt_obj = AiPrompt(client_id=client_id, prompt_text=prompt, is_active=True, added_at=datetime.today().date())
                db.add(prompt_obj)
                db.flush()
                
            for h, plat in platform_map.items():
                val = str(row.get(h) or "").strip().lower()
                if not val or val == "none" or val == "-":
                    continue
                    
                mentioned = val in ["true", "1", "yes", "y", "t"]
                
                stmt = select(AiMention).where(
                    AiMention.client_id == client_id,
                    AiMention.month == month,
                    AiMention.prompt_id == prompt_obj.id,
                    AiMention.platform == plat
                )
                existing = db.execute(stmt).scalar_one_or_none()
                
                if existing:
                    existing.mentioned = mentioned
                    existing.captured_on = datetime.today().date()
                else:
                    mentions_to_insert.append({
                        "client_id": client_id,
                        "month": month,
                        "prompt_id": prompt_obj.id,
                        "platform": plat,
                        "captured_on": datetime.today().date(),
                        "mentioned": mentioned,
                        "cited_pages": None,
                        "source": AiMentionSource.manual,
                        "raw_response": None
                    })
                
            success_count += 1
        except Exception as e:
            errors.append({"row": row_num, "error": str(e)})
            
    if mentions_to_insert:
        stmt = insert(AiMention).values(mentions_to_insert)
        stmt = stmt.on_conflict_do_update(
            constraint="uq_aimention_client_month_prompt_plat",
            set_={
                "mentioned": stmt.excluded.mentioned,
                "captured_on": stmt.excluded.captured_on
            }
        )
        db.execute(stmt)
    
    try:
        db.commit()
    except Exception as e:
        db.rollback()
        raise HTTPException(status_code=400, detail=f"Database error: {str(e)}")
        
    return {"status": "success", "rows_processed": success_count, "errors": errors}
