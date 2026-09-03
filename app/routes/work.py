from __future__ import annotations
import os
import uuid
import shutil
import datetime
from fastapi import APIRouter, Depends, HTTPException, status, UploadFile, File, Form
from pydantic import BaseModel
from sqlalchemy import select
from sqlalchemy.dialects.postgresql import insert
from sqlalchemy.orm import Session
from typing import List, Optional

from app.database import get_db
from app.models.activity import Activity
from app.models.screenshot import Screenshot
from app.models.report_snapshot import ReportSnapshot
from app.models.enums import UserRole, ReportStatus
from app.dependencies import RequireRole

router = APIRouter(
    prefix="/clients/{client_id}/work",
    tags=["work"],
    dependencies=[Depends(RequireRole([UserRole.super_admin, UserRole.manager, UserRole.user]))]
)

# Directory for screenshots
SCREENSHOTS_DIR = os.path.join(os.getcwd(), "app", "static", "screenshots")
os.makedirs(SCREENSHOTS_DIR, exist_ok=True)

class ActivityCreate(BaseModel):
    month: datetime.date
    activity_type: str
    count: int
    notes: str | None = None

class ActivityResponse(BaseModel):
    id: uuid.UUID
    month: datetime.date
    activity_type: str
    count: int
    notes: str | None

class ScreenshotResponse(BaseModel):
    id: uuid.UUID
    month: datetime.date
    keyword_id: uuid.UUID | None
    file_url: str
    caption: str | None

class WorkMonthResponse(BaseModel):
    activities: List[ActivityResponse]
    screenshots: List[ScreenshotResponse]
    next_month_plan: dict | None

class PlanUpdate(BaseModel):
    next_month_plan: dict | None

@router.get("")
def get_all_activities(
    client_id: uuid.UUID,
    page: int = 1,
    page_size: int = 25,
    start_date: datetime.date | None = None,
    end_date: datetime.date | None = None,
    db: Session = Depends(get_db)
):
    from sqlalchemy import func
    query = select(Activity).where(Activity.client_id == client_id)
    if start_date:
        query = query.where(Activity.month >= start_date)
    if end_date:
        query = query.where(Activity.month <= end_date)
        
    total = db.execute(select(func.count()).select_from(query.subquery())).scalar() or 0
    activities = db.execute(query.order_by(Activity.month.desc()).offset((page - 1) * page_size).limit(page_size)).scalars().all()
    return {"items": activities, "total": total, "page": page, "page_size": page_size}

@router.get("/{month_str}", response_model=WorkMonthResponse)
def get_work_done(client_id: uuid.UUID, month_str: str, db: Session = Depends(get_db)):
    """Fetch activities, screenshots, and the next_month_plan for the period."""
    try:
        target_month = datetime.datetime.strptime(month_str, "%Y-%m").date()
    except ValueError:
        raise HTTPException(status_code=400, detail="Invalid month format, expected YYYY-MM")

    # Activities
    activities = db.execute(
        select(Activity).where(
            Activity.client_id == client_id,
            Activity.month == target_month
        ).order_by(Activity.activity_type)
    ).scalars().all()

    # Screenshots
    screenshots = db.execute(
        select(Screenshot).where(
            Screenshot.client_id == client_id,
            Screenshot.month == target_month
        )
    ).scalars().all()

    # Next Month Plan from ReportSnapshot (find latest snapshot ending in this month)
    import calendar
    _, last_day = calendar.monthrange(target_month.year, target_month.month)
    month_end = target_month.replace(day=last_day)
    
    report = db.execute(
        select(ReportSnapshot).where(
            ReportSnapshot.client_id == client_id,
            ReportSnapshot.end_date >= target_month,
            ReportSnapshot.end_date <= month_end
        ).order_by(ReportSnapshot.end_date.desc())
    ).scalar_one_or_none()

    next_month_plan = report.next_month_plan if report else None

    return {
        "activities": activities,
        "screenshots": screenshots,
        "next_month_plan": next_month_plan
    }

@router.post("/activities", response_model=ActivityResponse, status_code=status.HTTP_201_CREATED)
def create_activity(client_id: uuid.UUID, payload: ActivityCreate, db: Session = Depends(get_db)):
    """Create a manual activity entry."""
    stmt = insert(Activity).values(
        client_id=client_id,
        month=payload.month,
        activity_type=payload.activity_type,
        count=payload.count,
        notes=payload.notes
    )
    stmt = stmt.on_conflict_do_update(
        constraint="uq_activity_client_month_type",
        set_={
            "count": stmt.excluded.count,
            "notes": stmt.excluded.notes
        }
    )
    result = db.execute(stmt.returning(Activity))
    activity = result.scalar_one()
    db.commit()
    return activity

@router.post("/screenshots", response_model=ScreenshotResponse, status_code=status.HTTP_201_CREATED)
def upload_screenshot(
    client_id: uuid.UUID,
    month: str = Form(...),
    caption: str | None = Form(None),
    keyword_id: str | None = Form(None),
    file: UploadFile = File(...),
    db: Session = Depends(get_db)
):
    """Upload a screenshot file."""
    try:
        target_month = datetime.datetime.strptime(month, "%Y-%m-%d").date()
    except ValueError:
        try:
            target_month = datetime.datetime.strptime(month, "%Y-%m").date()
        except ValueError:
            raise HTTPException(status_code=400, detail="Invalid month format, expected YYYY-MM or YYYY-MM-DD")

    file_extension = os.path.splitext(file.filename)[1] if file.filename else ".png"
    unique_filename = f"{uuid.uuid4()}{file_extension}"
    file_path = os.path.join(SCREENSHOTS_DIR, unique_filename)

    with open(file_path, "wb") as buffer:
        shutil.copyfileobj(file.file, buffer)

    # Relative URL to serve via StaticFiles
    file_url = f"/static/screenshots/{unique_filename}"

    parsed_kw_id = None
    if keyword_id and keyword_id.strip():
        try:
            parsed_kw_id = uuid.UUID(keyword_id.strip())
        except ValueError:
            raise HTTPException(status_code=400, detail="Invalid keyword_id format")

    stmt = insert(Screenshot).values(
        client_id=client_id,
        month=target_month,
        keyword_id=parsed_kw_id,
        file_url=file_url,
        caption=caption
    )
    stmt = stmt.on_conflict_do_update(
        constraint="uq_screenshot_client_month_kw",
        set_={
            "file_url": stmt.excluded.file_url,
            "caption": stmt.excluded.caption
        }
    )
    result = db.execute(stmt.returning(Screenshot))
    screenshot = result.scalar_one()
    db.commit()
    return screenshot

@router.delete("/screenshots/{screenshot_id}")
def delete_screenshot(
    client_id: uuid.UUID,
    screenshot_id: uuid.UUID,
    db: Session = Depends(get_db)
):
    """Delete a screenshot record and its associated file."""
    screenshot = db.execute(
        select(Screenshot).where(Screenshot.client_id == client_id, Screenshot.id == screenshot_id)
    ).scalar_one_or_none()
    
    if not screenshot:
        raise HTTPException(status_code=404, detail="Screenshot not found")
        
    if screenshot.file_url:
        filename = screenshot.file_url.split("/")[-1]
        file_path = os.path.join(SCREENSHOTS_DIR, filename)
        if os.path.exists(file_path):
            os.remove(file_path)
            
    db.delete(screenshot)
    db.commit()
    return {"status": "success"}


@router.put("/{month_str}/plan")
def update_plan(client_id: uuid.UUID, month_str: str, payload: PlanUpdate, db: Session = Depends(get_db)):
    """Edit the next_month_plan of a draft/review report."""
    target_month = datetime.date.fromisoformat(f"{month_str}-01")
    
    # Find the snapshot for the month
    import calendar
    _, last_day = calendar.monthrange(target_month.year, target_month.month)
    month_end = target_month.replace(day=last_day)
    
    report = db.execute(
        select(ReportSnapshot).where(
            ReportSnapshot.client_id == client_id,
            ReportSnapshot.end_date >= target_month,
            ReportSnapshot.end_date <= month_end
        ).order_by(ReportSnapshot.end_date.desc())
    ).scalar_one_or_none()
    
    if not report:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Report not found.")
        
    if report.status == ReportStatus.published:
        raise HTTPException(status.HTTP_409_CONFLICT, "Cannot edit a published report.")
        
    report.next_month_plan = payload.next_month_plan
    db.commit()
    
    return {"status": "success", "next_month_plan": report.next_month_plan}


import csv
import io

@router.post("/upload_csv", status_code=201)
def upload_work_csv(
    client_id: uuid.UUID,
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
) -> dict:
    if not (file.filename.endswith(".csv") or file.filename.endswith(".xlsx")):
        raise HTTPException(status_code=400, detail="Must be a CSV or Excel file")
    
    rows = []
    if file.filename.endswith(".xlsx"):
        import openpyxl
        content = file.file.read()
        wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
        ws = wb.active
        if not ws or ws.max_row < 1:
            raise HTTPException(status_code=400, detail="Empty Excel file")
        raw_headers = [str(c.value).strip() if c.value else "" for c in ws[1]]
        for row in ws.iter_rows(min_row=2, values_only=True):
            row_dict = {}
            for k, v in zip(raw_headers, row):
                row_dict[k] = v
            rows.append(row_dict)
    else:
        content = file.file.read().decode("utf-8")
        reader = csv.reader(io.StringIO(content))
        header_row = next(reader, None)
        if not header_row:
            raise HTTPException(status_code=400, detail="Empty CSV")
        raw_headers = [c.strip() for c in header_row]
        for row in reader:
            row_dict = {}
            for k, v in zip(raw_headers, row):
                row_dict[k] = v
            rows.append(row_dict)
            
    required_cols = {"activity_type", "count", "notes"}
    actual_cols = {col.lower() for col in raw_headers}
        
    work_to_insert = []
    success_count = 0
    errors = []
    
    for row_num, row in enumerate(rows, start=2):
        try:
            row_norm = {k.lower().replace(" ", "_"): v for k, v in row.items() if k}
            month = datetime.datetime.today().date().replace(day=1)
            activity_type = str(row_norm.get("activity_type") or "").strip()
            count_str = str(row_norm.get("count") or "1").strip()
            count = int(float(count_str)) if count_str else 1
            notes = str(row_norm.get("notes") or "").strip()
            if not notes: notes = None
            existing = db.execute(
                select(Activity).where(
                    Activity.client_id == client_id,
                    Activity.month == month,
                    Activity.activity_type == activity_type
                )
            ).scalar_one_or_none()

            if existing:
                existing.count = count
                existing.notes = notes
            else:
                db.add(
                    Activity(
                        client_id=client_id,
                        month=month,
                        activity_type=activity_type,
                        count=count,
                        notes=notes
                    )
                )
            db.flush()
            success_count += 1
        except Exception as e:
            errors.append({"row": row_num, "error": str(e)})
    
    try:
        db.commit()
    except Exception as e:
        db.rollback()
        raise HTTPException(status_code=400, detail=f"Database error: {str(e)}")
        
    return {"status": "success", "rows_processed": success_count, "errors": errors}

