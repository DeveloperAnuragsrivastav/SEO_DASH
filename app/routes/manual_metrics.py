from __future__ import annotations
"""Routes for Manual Metrics entry."""

import uuid
from datetime import date

from fastapi import APIRouter, Depends, HTTPException
from pydantic import BaseModel, Field
from sqlalchemy.orm import Session
from sqlalchemy import select

from app.database import get_db
from app.models.enums import MetricSource
from app.models.metric import Metric

from app.dependencies import RequireRole
from app.models.enums import UserRole

router = APIRouter(
    prefix="/clients/{client_id}",
    tags=["manual_metrics"],
    dependencies=[Depends(RequireRole([UserRole.super_admin, UserRole.manager, UserRole.user]))]
)


class ManualGSCRecord(BaseModel):
    captured_on: date
    clicks: int = Field(ge=0)
    impressions: int = Field(ge=0)
    ctr: float = Field(ge=0.0, le=1.0)
    position: float = Field(gt=0.0)
    dimension_key: str | None = None
    dimension_value: str | None = None


class ManualGSCInput(BaseModel):
    records: list[ManualGSCRecord]


@router.post("/manual-gsc", status_code=201)
def create_manual_gsc_metrics(
    client_id: uuid.UUID,
    data: ManualGSCInput,
    db: Session = Depends(get_db),
) -> dict:
    """Manually enter GSC metrics for a client."""

    metrics_to_insert = []

    for record in data.records:
        for metric_key, val in [
            ("clicks", float(record.clicks)),
            ("impressions", float(record.impressions)),
            ("ctr", record.ctr),
            ("position", record.position),
        ]:
            metrics_to_insert.append(
                Metric(
                    client_id=client_id,
                    provider="gsc",
                    metric_key=metric_key,
                    dimension_key=record.dimension_key,
                    dimension_value=record.dimension_value,
                    captured_on=record.captured_on,
                    value=val,
                    source=MetricSource.manual,
                )
            )

    if metrics_to_insert:
        db.add_all(metrics_to_insert)
        try:
            db.commit()
        except Exception as e:
            db.rollback()
            raise HTTPException(status_code=400, detail=f"Database error: {str(e)}")

    return {"status": "success", "rows_inserted": len(metrics_to_insert)}


class ManualGA4Record(BaseModel):
    captured_on: date
    sessions: int = Field(ge=0)
    users: int = Field(ge=0)
    engaged_sessions: int = Field(ge=0)
    conversions: int = Field(ge=0)
    revenue: float = Field(ge=0.0)
    dimension_key: str | None = None
    dimension_value: str | None = None


class ManualGA4Input(BaseModel):
    records: list[ManualGA4Record]


@router.post("/manual-ga4", status_code=201)
def create_manual_ga4_metrics(
    client_id: uuid.UUID,
    data: ManualGA4Input,
    db: Session = Depends(get_db),
) -> dict:
    """Manually enter GA4 metrics for a client."""

    metrics_to_insert = []

    for record in data.records:
        for metric_key, val in [
            ("sessions", float(record.sessions)),
            ("users", float(record.users)),
            ("engaged_sessions", float(record.engaged_sessions)),
            ("conversions", float(record.conversions)),
            ("revenue", record.revenue),
        ]:
            metrics_to_insert.append(
                Metric(
                    client_id=client_id,
                    provider="ga4",
                    metric_key=metric_key,
                    dimension_key=record.dimension_key,
                    dimension_value=record.dimension_value,
                    captured_on=record.captured_on,
                    value=val,
                    source=MetricSource.manual,
                )
            )

    if metrics_to_insert:
        db.add_all(metrics_to_insert)
        try:
            db.commit()
        except Exception as e:
            db.rollback()
            raise HTTPException(status_code=400, detail=f"Database error: {str(e)}")

    return {"status": "success", "rows_inserted": len(metrics_to_insert)}


class ManualGBPRecord(BaseModel):
    captured_on: date
    impressions_desktop_maps: int = Field(ge=0)
    impressions_desktop_search: int = Field(ge=0)
    impressions_mobile_maps: int = Field(ge=0)
    impressions_mobile_search: int = Field(ge=0)
    calls: int = Field(ge=0)
    direction_requests: int = Field(ge=0)
    website_clicks: int = Field(ge=0)
    bookings: int = Field(ge=0)


class ManualGBPInput(BaseModel):
    records: list[ManualGBPRecord]


@router.post("/manual-gbp", status_code=201)
def create_manual_gbp_metrics(
    client_id: uuid.UUID,
    data: ManualGBPInput,
    db: Session = Depends(get_db),
) -> dict:
    """Manually enter GBP metrics for a client."""

    metrics_to_insert = []

    for record in data.records:
        for metric_key, val in [
            ("impressions_desktop_maps", float(record.impressions_desktop_maps)),
            ("impressions_desktop_search", float(record.impressions_desktop_search)),
            ("impressions_mobile_maps", float(record.impressions_mobile_maps)),
            ("impressions_mobile_search", float(record.impressions_mobile_search)),
            ("calls", float(record.calls)),
            ("direction_requests", float(record.direction_requests)),
            ("website_clicks", float(record.website_clicks)),
            ("bookings", float(record.bookings)),
        ]:
            metrics_to_insert.append(
                Metric(
                    client_id=client_id,
                    provider="gbp",
                    metric_key=metric_key,
                    dimension_key="",
                    dimension_value="",
                    captured_on=record.captured_on,
                    value=val,
                    source=MetricSource.manual,
                )
            )

    if metrics_to_insert:
        db.add_all(metrics_to_insert)
        try:
            db.commit()
        except Exception as e:
            db.rollback()
            raise HTTPException(status_code=400, detail=f"Database error: {str(e)}")

    return {"status": "success", "rows_inserted": len(metrics_to_insert)}

@router.get("/manual-gbp")
def get_manual_gbp_metrics(
    client_id: uuid.UUID,
    page: int = 1,
    page_size: int = 25,
    db: Session = Depends(get_db),
) -> dict:
    from sqlalchemy import select
    from collections import defaultdict
    
    metrics = db.execute(
        select(Metric).where(
            Metric.client_id == client_id,
            Metric.provider == "gbp",
            Metric.source == MetricSource.manual
        )
    ).scalars().all()
    
    # Group by captured_on
    grouped = defaultdict(dict)
    for m in metrics:
        grouped[m.captured_on.isoformat()][m.metric_key] = m.value
        
    records = []
    for date_str, values in grouped.items():
        records.append({
            "captured_on": date_str,
            "impressions_desktop_maps": values.get("impressions_desktop_maps", 0),
            "impressions_desktop_search": values.get("impressions_desktop_search", 0),
            "impressions_mobile_maps": values.get("impressions_mobile_maps", 0),
            "impressions_mobile_search": values.get("impressions_mobile_search", 0),
            "calls": values.get("calls", 0),
            "direction_requests": values.get("direction_requests", 0),
            "website_clicks": values.get("website_clicks", 0),
            "bookings": values.get("bookings", 0),
        })
        
    sorted_records = sorted(records, key=lambda x: x["captured_on"], reverse=True)
    
    total = len(sorted_records)
    start_idx = (page - 1) * page_size
    end_idx = start_idx + page_size
    paginated = sorted_records[start_idx:end_idx]
    
    return {"items": paginated, "total": total, "page": page, "page_size": page_size}


import csv
import io
from fastapi import File, UploadFile
from datetime import datetime

@router.post("/manual-gsc/upload_csv", status_code=201)
def upload_manual_gsc_csv(
    client_id: uuid.UUID,
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
) -> dict:
    if not file.filename.endswith(".csv"):
        raise HTTPException(status_code=400, detail="Must be a CSV file")
    
    content = file.file.read().decode("utf-8")
    reader = csv.DictReader(io.StringIO(content))
    if not reader.fieldnames:
        raise HTTPException(status_code=400, detail="Empty CSV")
        
    required_cols = {"date", "clicks", "impressions", "ctr", "position"}
    actual_cols = {col.lower() for col in reader.fieldnames}
    if not required_cols.issubset(actual_cols):
        raise HTTPException(status_code=400, detail=f"CSV must contain columns: {', '.join(required_cols)}")
        
    metrics_to_insert = []
    success_count = 0
    errors = []
    
    for row_num, row in enumerate(reader, start=2):
        try:
            date_str = row.get("date", "").strip()
            if not date_str:
                continue
            captured_on = None
            for fmt in ("%Y-%m-%d", "%m/%d/%Y", "%m/%d/%y", "%Y/%m/%d"):
                try:
                    captured_on = datetime.strptime(date_str.split(" ")[0].split("T")[0], fmt).date()
                    break
                except ValueError:
                    pass
            if not captured_on: raise ValueError(f"Invalid date {date_str}")
            
            for metric_key in ["clicks", "impressions", "ctr", "position"]:
                val = float(row.get(metric_key, 0))
                
                # Check for existing metric
                stmt = select(Metric).where(
                    Metric.client_id == client_id,
                    Metric.provider == "gsc",
                    Metric.metric_key == metric_key,
                    Metric.captured_on == captured_on,
                    Metric.source == MetricSource.manual
                )
                existing = db.execute(stmt).scalar_one_or_none()
                
                if existing:
                    existing.value = val
                else:
                    metrics_to_insert.append(
                        Metric(
                            client_id=client_id,
                            provider="gsc",
                            metric_key=metric_key,
                            captured_on=captured_on,
                            value=val,
                            source=MetricSource.manual,
                        )
                    )
            success_count += 1
        except Exception as e:
            errors.append({"row": row_num, "error": str(e)})
            
    if metrics_to_insert:
        db.add_all(metrics_to_insert)
    
    try:
        db.commit()
    except Exception as e:
        db.rollback()
        raise HTTPException(status_code=400, detail=f"Database error: {str(e)}")
        
    return {"status": "success", "rows_processed": success_count, "errors": errors}


@router.post("/manual-ga4/upload_csv", status_code=201)
def upload_manual_ga4_csv(
    client_id: uuid.UUID,
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
) -> dict:
    if not file.filename.endswith(".csv"):
        raise HTTPException(status_code=400, detail="Must be a CSV file")
    
    content = file.file.read().decode("utf-8")
    reader = csv.DictReader(io.StringIO(content))
    if not reader.fieldnames:
        raise HTTPException(status_code=400, detail="Empty CSV")
        
    required_cols = {"date", "sessions", "users", "engaged_sessions", "conversions", "revenue"}
    actual_cols = {col.lower() for col in reader.fieldnames}
    if not required_cols.issubset(actual_cols):
        raise HTTPException(status_code=400, detail=f"CSV must contain columns: {', '.join(required_cols)}")
        
    metrics_to_insert = []
    success_count = 0
    errors = []
    
    for row_num, row in enumerate(reader, start=2):
        try:
            date_str = row.get("date", "").strip()
            if not date_str:
                continue
            captured_on = None
            for fmt in ("%Y-%m-%d", "%m/%d/%Y", "%m/%d/%y", "%Y/%m/%d"):
                try:
                    captured_on = datetime.strptime(date_str.split(" ")[0].split("T")[0], fmt).date()
                    break
                except ValueError:
                    pass
            if not captured_on: raise ValueError(f"Invalid date {date_str}")
            
            for metric_key in ["sessions", "users", "engaged_sessions", "conversions", "revenue"]:
                val = float(row.get(metric_key, 0))
                
                stmt = select(Metric).where(
                    Metric.client_id == client_id,
                    Metric.provider == "ga4",
                    Metric.metric_key == metric_key,
                    Metric.captured_on == captured_on,
                    Metric.source == MetricSource.manual
                )
                existing = db.execute(stmt).scalar_one_or_none()
                
                if existing:
                    existing.value = val
                else:
                    metrics_to_insert.append(
                        Metric(
                            client_id=client_id,
                            provider="ga4",
                            metric_key=metric_key,
                            captured_on=captured_on,
                            value=val,
                            source=MetricSource.manual,
                        )
                    )
            success_count += 1
        except Exception as e:
            errors.append({"row": row_num, "error": str(e)})
            
    if metrics_to_insert:
        db.add_all(metrics_to_insert)
    
    try:
        db.commit()
    except Exception as e:
        db.rollback()
        raise HTTPException(status_code=400, detail=f"Database error: {str(e)}")
        
    return {"status": "success", "rows_processed": success_count, "errors": errors}


@router.post("/manual-gbp/upload_csv", status_code=201)
def upload_manual_gbp_csv(
    client_id: uuid.UUID,
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
) -> dict:
    if not (file.filename.endswith(".csv") or file.filename.endswith(".xlsx")):
        raise HTTPException(status_code=400, detail="Must be a CSV or Excel file")
    
    rows = []
    if file.filename.endswith(".xlsx"):
        import openpyxl
        import io
        content = file.file.read()
        wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
        ws = wb.active
        if not ws or ws.max_row < 1:
            raise HTTPException(status_code=400, detail="Empty Excel file")
        raw_headers = [str(c.value).strip() if c.value else "" for c in ws[1]]
        for row in ws.iter_rows(min_row=2, values_only=True):
            row_dict = {}
            for k, v in zip(raw_headers, row):
                row_dict[k] = v
            rows.append(row_dict)
    else:
        import io
        content = file.file.read().decode("utf-8")
        reader = csv.reader(io.StringIO(content))
        header_row = next(reader, None)
        if not header_row:
            raise HTTPException(status_code=400, detail="Empty CSV")
        raw_headers = [c.strip() for c in header_row]
        for row in reader:
            row_dict = {}
            for k, v in zip(raw_headers, row):
                row_dict[k] = v
            rows.append(row_dict)
            
    required_cols = {"date", "impressions_desktop_maps", "impressions_desktop_search", "impressions_mobile_maps", "impressions_mobile_search", "calls", "direction_requests", "website_clicks", "bookings"}
    actual_cols = {col.lower().replace(" ", "_") for col in raw_headers}
    if not required_cols.issubset(actual_cols):
        raise HTTPException(status_code=400, detail=f"CSV must contain columns: {', '.join(required_cols)}")
        
    metrics_to_insert = []
    success_count = 0
    errors = []
    
    for row_num, row in enumerate(rows, start=2):
        try:
            row_norm = {k.lower().replace(" ", "_"): v for k, v in row.items() if k}
            date_str = str(row_norm.get("date") or "").strip()
            if not date_str:
                continue
            captured_on = None
            for fmt in ("%b'%y", "%B %Y", "%Y-%m-%d", "%m/%d/%Y", "%m/%d/%y", "%Y/%m/%d"):
                try:
                    clean_str = date_str.split(" ")[0].split("T")[0] if 'T' in date_str else date_str
                    captured_on = datetime.strptime(clean_str, fmt).date()
                    break
                except ValueError:
                    pass
            if not captured_on: raise ValueError(f"Invalid date {date_str}")
            
            for metric_key in ["impressions_desktop_maps", "impressions_desktop_search", "impressions_mobile_maps", "impressions_mobile_search", "calls", "direction_requests", "website_clicks", "bookings"]:
                val = float(row_norm.get(metric_key) or 0)
                
                stmt = select(Metric).where(
                    Metric.client_id == client_id,
                    Metric.provider == "gbp",
                    Metric.metric_key == metric_key,
                    Metric.dimension_key == "",
                    Metric.dimension_value == "",
                    Metric.captured_on == captured_on,
                    Metric.source == MetricSource.manual
                )
                existing = db.execute(stmt).scalar_one_or_none()
                
                if existing:
                    existing.value = val
                else:
                    metrics_to_insert.append(
                        Metric(
                            client_id=client_id,
                            provider="gbp",
                            metric_key=metric_key,
                            dimension_key="",
                            dimension_value="",
                            captured_on=captured_on,
                            value=val,
                            source=MetricSource.manual,
                        )
                    )
            success_count += 1
        except Exception as e:
            errors.append({"row": row_num, "error": str(e)})
            
    if metrics_to_insert:
        db.add_all(metrics_to_insert)
    
    try:
        db.commit()
    except Exception as e:
        db.rollback()
        raise HTTPException(status_code=400, detail=f"Database error: {str(e)}")
        
    return {"status": "success", "rows_processed": success_count, "errors": errors}
