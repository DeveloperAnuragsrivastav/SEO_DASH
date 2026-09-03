from __future__ import annotations
import uuid
import datetime
from calendar import monthrange
from fastapi import APIRouter, Depends, HTTPException, status
from pydantic import BaseModel, Field
from sqlalchemy import select, func
from sqlalchemy.dialects.postgresql import insert
from sqlalchemy.orm import Session
from typing import List, Optional

from app.database import get_db
from app.models.link import Link
from app.models.enums import LinkStatus, UserRole
from app.dependencies import RequireRole

router = APIRouter(
    prefix="/clients/{client_id}/links",
    tags=["links"],
    dependencies=[Depends(RequireRole([UserRole.super_admin, UserRole.manager, UserRole.user]))]
)

def get_month_boundaries(year: int, month: int) -> tuple[datetime.date, datetime.date]:
    _, last_day = monthrange(year, month)
    return datetime.date(year, month, 1), datetime.date(year, month, last_day)

class LinkCreate(BaseModel):
    domain: str
    url: str
    activity_type: str
    status: LinkStatus = LinkStatus.active
    dr: int | None = None
    created_on: datetime.date

class LinkResponse(BaseModel):
    id: uuid.UUID
    domain: str | None
    url: str | None
    activity_type: str
    count: int
    status: LinkStatus
    dr: int | None
    last_checked: datetime.datetime | None
    created_on: datetime.date

class LinksMonthResponse(BaseModel):
    kpis: dict
    items: List[LinkResponse]
    total: int
    page: int
    page_size: int

@router.get("")
def get_all_links(
    client_id: uuid.UUID,
    page: int = 1,
    page_size: int = 25,
    start_date: datetime.date | None = None,
    end_date: datetime.date | None = None,
    db: Session = Depends(get_db)
):
    query = select(Link).where(Link.client_id == client_id)
    if start_date:
        query = query.where(Link.month >= start_date)
    if end_date:
        query = query.where(Link.month <= end_date)
        
    total = db.execute(select(func.count()).select_from(query.subquery())).scalar() or 0
    links_db = db.execute(query.order_by(Link.created_on.desc()).offset((page - 1) * page_size).limit(page_size)).scalars().all()
    return {"items": links_db, "total": total, "page": page, "page_size": page_size}

@router.get("/{month_str}", response_model=LinksMonthResponse)
def get_links(client_id: uuid.UUID, month_str: str, page: int = 1, page_size: int = 25, db: Session = Depends(get_db)):
    """Fetch links for the client/period with KPI summaries."""
    try:
        target_month = datetime.datetime.strptime(month_str, "%Y-%m").date()
    except ValueError:
        raise HTTPException(status_code=400, detail="Invalid month format, expected YYYY-MM")

    start_date, end_date = get_month_boundaries(target_month.year, target_month.month)
    
    base_query = select(Link).where(
        Link.client_id == client_id,
        Link.created_on >= start_date,
        Link.created_on <= end_date
    )

    total = db.execute(select(func.count()).select_from(base_query.subquery())).scalar() or 0

    links_db = db.execute(
        base_query.order_by(Link.created_on.desc()).offset((page - 1) * page_size).limit(page_size)
    ).scalars().all()

    # Calculate KPIs (for KPIs, we might need all links for the month, not just paginated)
    all_links_for_kpi = db.execute(base_query).scalars().all()
    unique_domains = len(set(l.domain for l in all_links_for_kpi))
    activity_breakdown = {}
    for l in all_links_for_kpi:
        activity_breakdown[l.activity_type] = activity_breakdown.get(l.activity_type, 0) + 1

    kpis = {
        "total_links": len(all_links_for_kpi),
        "unique_domains": unique_domains,
        "activity_breakdown": activity_breakdown
    }

    return {
        "kpis": kpis,
        "items": links_db,
        "total": total,
        "page": page,
        "page_size": page_size
    }

@router.post("", response_model=LinkResponse, status_code=status.HTTP_201_CREATED)
def create_link(client_id: uuid.UUID, payload: LinkCreate, db: Session = Depends(get_db)):
    """Create a manual link entry."""
    stmt = insert(Link).values(
        client_id=client_id,
        created_on=payload.created_on,
        month=payload.created_on.replace(day=1),
        activity_type=payload.activity_type,
        domain=payload.domain,
        url=payload.url,
        status=payload.status,
        dr=payload.dr,
        last_checked=None
    )
    stmt = stmt.on_conflict_do_update(
        constraint="uq_link_client_month_url",
        set_={
            "activity_type": stmt.excluded.activity_type,
            "domain": stmt.excluded.domain,
            "status": stmt.excluded.status,
            "dr": stmt.excluded.dr
        }
    )
    result = db.execute(stmt.returning(Link))
    link = result.scalar_one()
    db.commit()
    return link


import csv
import io
from fastapi import File, UploadFile

@router.post("/upload_csv", status_code=201)
def upload_links_csv(
    client_id: uuid.UUID,
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
) -> dict:
    if not (file.filename.endswith(".csv") or file.filename.endswith(".xlsx")):
        raise HTTPException(status_code=400, detail="Must be a CSV or Excel file")
    
    rows = []
    
    if file.filename.endswith(".xlsx"):
        import openpyxl
        content = file.file.read()
        wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
        ws = wb.active
        if not ws or ws.max_row < 1:
            raise HTTPException(status_code=400, detail="Empty Excel file")
        raw_headers = [str(c.value).strip() if c.value else "" for c in ws[1]]
        for row in ws.iter_rows(min_row=2, values_only=True):
            row_dict = {}
            for k, v in zip(raw_headers, row):
                row_dict[k] = v
            rows.append(row_dict)
    else:
        content = file.file.read().decode("utf-8")
        reader = csv.reader(io.StringIO(content))
        header_row = next(reader, None)
        if not header_row:
            raise HTTPException(status_code=400, detail="Empty CSV")
        raw_headers = [c.strip() for c in header_row]
        for row in reader:
            row_dict = {}
            for k, v in zip(raw_headers, row):
                row_dict[k] = v
            rows.append(row_dict)
        
    links_to_insert = []
    success_count = 0
    errors = []
    
    import urllib.parse
    
    for row_num, row in enumerate(rows, start=2):
        row_norm = {k.lower().replace(" ", "_"): v for k, v in row.items() if k}
        try:
            month_str = str(row_norm.get("month") or row_norm.get("date") or "").strip()
            activity_type = str(row_norm.get("activity_name") or row_norm.get("activity_type") or "").strip()
            count_str = str(row_norm.get("count") or "1").strip()
            url = str(row_norm.get("url") or "").strip()
            
            if not activity_type:
                continue
                
            count = int(float(count_str)) if count_str else 1
            
            created_on = None
            if month_str:
                for fmt in ("%b'%y", "%B %Y", "%Y-%m-%d", "%m/%d/%Y", "%m/%d/%y", "%Y/%m/%d"):
                    try:
                        # Clean up formatting for parsing (e.g. Aug'26)
                        clean_str = month_str.split(" ")[0].split("T")[0] if 'T' in month_str else month_str
                        created_on = datetime.datetime.strptime(clean_str, fmt).date()
                        break
                    except ValueError:
                        pass
            if not created_on:
                created_on = datetime.datetime.today().date().replace(day=1)
                
            domain = None
            if url:
                parsed_uri = urllib.parse.urlparse(url)
                domain = '{uri.netloc}'.format(uri=parsed_uri).replace("www.", "")
                
            links_to_insert.append({
                "client_id": client_id,
                "month": created_on.replace(day=1),
                "created_on": created_on,
                "activity_type": activity_type,
                "domain": domain if domain else None,
                "url": url if url else None,
                "count": count,
                "status": LinkStatus.active,
                "dr": None,
                "last_checked": None
            })
            success_count += 1
        except Exception as e:
            errors.append({"row": row_num, "error": str(e)})
            
    if links_to_insert:
        stmt = insert(Link).values(links_to_insert)
        stmt = stmt.on_conflict_do_update(
            constraint="uq_link_client_month_url_activity",
            set_={
                "count": stmt.excluded.count,
                "status": stmt.excluded.status
            }
        )
        db.execute(stmt)
    
    try:
        db.commit()
    except Exception as e:
        db.rollback()
        raise HTTPException(status_code=400, detail=f"Database error: {str(e)}")
        
    return {"status": "success", "rows_processed": success_count, "errors": errors}

